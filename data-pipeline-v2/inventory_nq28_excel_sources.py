#!/usr/bin/env python3
"""Inventory and validate the original NQ28 Excel source pack.

The script is intentionally fail-closed. It never edits workbooks. It validates that
exactly one numbered workbook exists for each legal area 1..95, records hashes and
sheet-level structural metadata, and emits a deterministic manifest for later
parsing and audit.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

AREA_PREFIX_RE = re.compile(r"^\s*(\d{1,3})(?:\s*[.\-_)]|\s+)")
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def workbook_metadata(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            non_empty = 0
            formula_cells = 0
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None or value == "":
                        continue
                    non_empty += 1
                    if isinstance(value, str) and value.startswith("="):
                        formula_cells += 1
            sheets.append(
                {
                    "title": worksheet.title,
                    "state": worksheet.sheet_state,
                    "maxRow": worksheet.max_row,
                    "maxColumn": worksheet.max_column,
                    "nonEmptyCells": non_empty,
                    "formulaCells": formula_cells,
                }
            )
    finally:
        workbook.close()
    return {"sheetCount": len(sheets), "sheets": sheets}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--expected-areas", type=int, default=95)
    args = parser.parse_args()

    source_dir = Path(args.source_dir)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    files = sorted(
        path
        for path in source_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    ) if source_dir.is_dir() else []

    entries: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    by_area: dict[int, list[str]] = {}

    for path in files:
        relative = path.relative_to(source_dir).as_posix()
        match = AREA_PREFIX_RE.match(path.name)
        area_no = int(match.group(1)) if match else None
        if area_no is None:
            errors.append({"code": "missing-area-prefix", "file": relative})
        elif not 1 <= area_no <= args.expected_areas:
            errors.append({"code": "area-prefix-out-of-range", "file": relative, "area": area_no})
        else:
            by_area.setdefault(area_no, []).append(relative)

        try:
            metadata = workbook_metadata(path)
            workbook_error = None
        except Exception as exc:  # preserve evidence; do not guess around damaged files
            metadata = {"sheetCount": 0, "sheets": []}
            workbook_error = f"{type(exc).__name__}: {exc}"
            errors.append({"code": "workbook-open-failed", "file": relative, "detail": workbook_error})

        entries.append(
            {
                "area": area_no,
                "file": relative,
                "sizeBytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "workbook": metadata,
                "error": workbook_error,
            }
        )

    expected = set(range(1, args.expected_areas + 1))
    present = set(by_area)
    missing = sorted(expected - present)
    duplicates = {str(area): names for area, names in sorted(by_area.items()) if len(names) != 1}

    if len(files) != args.expected_areas:
        errors.append({"code": "unexpected-workbook-count", "expected": args.expected_areas, "actual": len(files)})
    if missing:
        errors.append({"code": "missing-areas", "areas": missing})
    if duplicates:
        errors.append({"code": "duplicate-area-files", "areas": duplicates})

    manifest = {
        "schemaVersion": 1,
        "sourceKind": "nq28-original-excel-pack",
        "sourceDirectory": source_dir.as_posix(),
        "expectedAreaCount": args.expected_areas,
        "workbookCount": len(files),
        "coveredAreaCount": len(present & expected),
        "missingAreas": missing,
        "duplicateAreas": duplicates,
        "status": "passed" if not errors else "blocked",
        "errors": errors,
        "files": entries,
    }
    output_path.write_text(
        json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({key: manifest[key] for key in ("status", "workbookCount", "coveredAreaCount", "missingAreas", "duplicateAreas")}, ensure_ascii=False, indent=2))
    return 0 if not errors else 2


if __name__ == "__main__":
    raise SystemExit(main())
